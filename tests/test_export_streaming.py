"""Потоковый экспорт: потолок строк, пометка об обрезке, чтение курсором без OFFSET."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event
from tests.test_api_health import client as client_fixture  # noqa: F401

from backend.settings import get_settings


@pytest.fixture
def client(client_fixture, monkeypatch):  # noqa: F811
    monkeypatch.setenv("EXPORT_MAX_ROWS", "2")
    get_settings.cache_clear()
    yield client_fixture
    get_settings.cache_clear()


def _seed(n: int = 3) -> uuid.UUID:
    from backend.db.repository import create_parse_run, upsert_vacancies
    from backend.db.session import session_scope
    from tg_jobs_parser.models.vacancy import Vacancy

    with session_scope() as session:
        run_id = create_parse_run(session, config_snapshot={}).id
        vacancies = [
            Vacancy(
                id=f"exp-{i}",
                source_post_key=f"chan-{i}",
                source_channel="chan",
                source_url="https://t.me/s/chan",
                post_url=f"https://t.me/chan/{i}",
                post_date=datetime(2026, 9, 1 + i, tzinfo=UTC),
                raw_text=f"Python dev {i}\x01",  # управляющий символ Excel не принимает
                cleaned_text=f"Python dev {i}",
                detected_specializations=["Python"],
            )
            for i in range(n)
        ]
        upsert_vacancies(session, vacancies, parse_run_id=run_id)
    return run_id


def _sheets(content: bytes):
    wb = load_workbook(io.BytesIO(content))
    rows = list(wb["vacancies"].iter_rows(values_only=True))
    summary = dict(wb["metrics_summary"].iter_rows(min_row=2, values_only=True))
    return rows[0], rows[1:], summary


def test_vacancies_export_is_capped_and_marked(client: TestClient):
    _seed()
    r = client.get("/api/v1/vacancies/export.xlsx")

    assert r.status_code == 200, r.text[:300]
    assert r.headers["X-Export-Truncated"] == "2/3"
    header, rows, summary = _sheets(r.content)
    post_urls = [row[header.index("post_url")] for row in rows]
    assert post_urls == ["https://t.me/chan/2", "https://t.me/chan/1"]  # post_date desc, первые 2
    assert summary["exported_rows"] == 2 and summary["matching_rows"] == 3
    assert "обрезано до 2 из 3" in summary["truncated"]
    assert rows[0][header.index("raw_text")] == "Python dev 2"
    assert isinstance(rows[0][header.index("post_date")], datetime)


def test_run_exports_are_capped(client: TestClient):
    run_id = _seed()

    xlsx = client.get(f"/api/v1/runs/{run_id}/export.xlsx")
    assert xlsx.headers["X-Export-Truncated"] == "2/3"
    _, rows, _ = _sheets(xlsx.content)
    assert len(rows) == 2

    js_response = client.get(f"/api/v1/runs/{run_id}/export.json")
    assert js_response.headers["X-Export-Truncated"] == "2/3"
    js = js_response.json()
    assert len(js["vacancies_unique"]) == 2
    assert js["export_truncated"]["total"] == 3


def test_export_under_cap_has_no_truncation_mark(client: TestClient):
    _seed(n=2)
    r = client.get("/api/v1/vacancies/export.xlsx")
    assert "X-Export-Truncated" not in r.headers
    _, rows, summary = _sheets(r.content)
    assert len(rows) == 2 and "truncated" not in summary


def test_export_iterator_reads_once_and_keeps_session_light(client: TestClient):
    """Один SELECT, строки не копятся в identity map сессии."""
    from backend.db.session import get_engine, session_scope
    from backend.db.vacancy_repository import iter_export_vacancies
    from backend.schemas import VacancyQuery

    _seed(n=3)
    selects: list[str] = []

    def _log(conn, cursor, statement, params, context, executemany):
        if "FROM vacancies" in statement:
            selects.append(statement)

    engine = get_engine()
    event.listen(engine, "before_cursor_execute", _log)
    try:
        with session_scope() as session:
            got = [v.id for v in iter_export_vacancies(session, VacancyQuery(), limit=10)]
            assert len(session.identity_map) == 0
    finally:
        event.remove(engine, "before_cursor_execute", _log)

    assert got == ["exp-2", "exp-1", "exp-0"]
    assert len(selects) == 1  # один запрос курсором, а не цикл OFFSET-страниц
