"""Экспорт вакансий в Excel (потоковая запись через openpyxl)."""

from __future__ import annotations

import re
from collections.abc import Iterable
from datetime import UTC, datetime
from typing import Any

from openpyxl.utils import get_column_letter

from tg_jobs_parser.models.metrics import MetricsSummary
from tg_jobs_parser.models.vacancy import Vacancy
from tg_jobs_parser.utils.text import post_id_display

_DATETIME_FMT = "DD.MM.YYYY HH:MM:SS"


def _naive_utc_for_excel(dt: datetime | None) -> datetime | None:
    """openpyxl не принимает timezone-aware datetime в ячейках Excel."""
    if dt is None:
        return None
    if getattr(dt, "tzinfo", None):
        return dt.astimezone(UTC).replace(tzinfo=None)
    return dt


def _excel_single_line(text: str | None) -> str:
    """Один сплошной ряд без переводов строк — строка таблицы не раздувается по высоте."""
    if not text:
        return ""
    t = str(text).replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"\s+", " ", t).strip()


def _vacancy_row_for_excel(v: Vacancy, *, include_parse_dedup_columns: bool = True) -> dict[str, Any]:
    """Колонки без префикса ``detected_``, читаемые имена для отчёта."""
    d = v.model_dump(mode="python")
    row: dict[str, Any] = {
        "post_id": post_id_display(str(d.get("source_post_key") or "")),
        "source_channel": d.get("source_channel"),
        "post_url": d.get("post_url"),
        "post_date": _naive_utc_for_excel(d.get("post_date")),
        "specializations": ", ".join(d.get("detected_specializations") or []),
        "grade": d.get("detected_grade"),
        "salary_min": d.get("detected_salary_min"),
        "salary_max": d.get("detected_salary_max"),
        "salary_avg": d.get("detected_salary_avg"),
        "salary_currency": d.get("salary_currency"),
        "work_format": d.get("work_format"),
        "country": d.get("country"),
        "city": d.get("city"),
        "company": d.get("company"),
        "contact": d.get("contact"),
        "application_urls": ", ".join(d.get("application_urls") or []),
        "extraction_source": d.get("extraction_source"),
        "ai_confidence": d.get("ai_confidence"),
        "ai_notes": d.get("ai_notes"),
        "technologies": ", ".join(d.get("technologies") or []),
        "required_experience": d.get("required_experience"),
        "employment_type": d.get("employment_type"),
        "vacancy_score": d.get("vacancy_score"),
    }
    if include_parse_dedup_columns:
        row["is_duplicate"] = d.get("is_duplicate")
        row["duplicate_group_id"] = d.get("duplicate_group_id")
        row["duplicate_reason"] = d.get("duplicate_reason")
        row["parsed_at"] = _naive_utc_for_excel(d.get("parsed_at"))
        row["source_url"] = d.get("source_url")
    else:
        row["duplicate_group_id"] = d.get("duplicate_group_id")
    row["raw_text"] = _excel_single_line(d.get("raw_text"))
    row["cleaned_text"] = _excel_single_line(d.get("cleaned_text"))
    return row


# Ширина колонок потокового листа: содержимое заранее не просмотреть (write_only).
_STREAM_WIDTHS: dict[str, float] = {
    "post_id": 14,
    "source_channel": 20,
    "post_url": 32,
    "post_date": 22,
    "specializations": 28,
    "company": 24,
    "contact": 24,
    "application_urls": 32,
    "ai_notes": 40,
    "technologies": 32,
    "duplicate_group_id": 20,
    "raw_text": 52,
    "cleaned_text": 52,
}
_STREAM_DEFAULT_WIDTH = 14.0
# Лимит длины ячейки Excel.
_EXCEL_CELL_MAX = 32_767


def _stream_cell(ws, value):
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)[:_EXCEL_CELL_MAX]
    cell = WriteOnlyCell(ws, value=value)
    if isinstance(value, datetime):
        cell.number_format = _DATETIME_FMT
    return cell


def write_vacancies_xlsx(
    fileobj,
    vacancies: Iterable[Vacancy],
    *,
    metrics: MetricsSummary | None = None,
    summary_extra: list[tuple[str, Any]] | None = None,
) -> int:
    """Excel для API потоком: лист вакансий + опционально сводка метрик. Возвращает число строк.

    openpyxl ``write_only``: строки уходят во временный файл по мере итерации, в памяти —
    одна строка, а не весь лист (как у pandas). ``summary_extra`` — дополнительные строки
    сводки (например, пометка об обрезке экспорта).
    """
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("vacancies")
    written = 0
    for v in vacancies:
        row = _vacancy_row_for_excel(v, include_parse_dedup_columns=False)
        if written == 0:
            for idx, header in enumerate(row, start=1):
                ws.column_dimensions[get_column_letter(idx)].width = _STREAM_WIDTHS.get(header, _STREAM_DEFAULT_WIDTH)
            ws.append(list(row))
        ws.append([_stream_cell(ws, value) for value in row.values()])
        written += 1

    summary_rows: list[tuple[str, Any]] = []
    if metrics is not None:
        summary_rows += [
            ("total_posts_fetched", metrics.total_posts_fetched),
            ("unique_vacancies", metrics.unique_vacancies),
            ("duplicate_count", metrics.duplicate_count),
            ("vacancies_with_salary", metrics.vacancies_with_salary),
            ("vacancies_without_salary", metrics.vacancies_without_salary),
        ]
    summary_rows += summary_extra or []
    if summary_rows:
        summary = wb.create_sheet("metrics_summary")
        summary.column_dimensions["A"].width = 28
        summary.column_dimensions["B"].width = 60
        summary.append(["metric", "value"])
        for name, value in summary_rows:
            summary.append([name, value])
    wb.save(fileobj)
    return written
